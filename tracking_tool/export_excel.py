"""Esporta results.json (o results_iso.json) in un Excel di sintesi."""
import argparse, json
import openpyxl
from openpyxl.styles import PatternFill, Font

ap = argparse.ArgumentParser()
ap.add_argument("--results", required=True)
ap.add_argument("--out", default="Verifica_coppie.xlsx")
args = ap.parse_args()
res = json.load(open(args.results))
cols = ["Signal Name", "Sw Variable", "status", "ref_1K136", "only_1K137", "frame_hex", "frame_cfg", "cycle_ms",
        "cmm_dir", "kind", "match_frac", "delay_samples", "remap_frac", "norm_err", "corr",
        "sig_varies", "sw_varies", "sig_min", "sig_max", "sw_min", "sw_max", "n_samples"]
COL = {"OK": "C6EFCE", "REMAP": "C6EFCE", "CONFIG_EXCLUDED": "D6E4F0", "ISO_1K136": "D9EAD3",
       "SUSPECT": "FFEB9C", "MISMATCH": "FFC7CE", "SW_NOT_EVOLVING": "FFC7CE"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Verifica coppie"
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True)
for r in res:
    ws.append([r.get(c, "") for c in cols])
    fill = COL.get(r["status"], "EDEDED")
    ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=fill)
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
for i, w in enumerate([30, 38, 18], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws2 = wb.create_sheet("Riepilogo")
from collections import Counter
for k, v in Counter(r["status"] for r in res).most_common():
    ws2.append([k, v])
wb.save(args.out); print("Scritto", args.out)
