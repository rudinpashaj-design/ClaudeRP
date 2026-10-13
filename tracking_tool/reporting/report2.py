import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

results = json.load(open(os.environ.get("RESULTS", "results_iso.json")))
dbc = json.load(open(os.environ.get("DBC_LOOKUP", "dbc_lookup.json")))
MEAS = os.environ.get("MEAS_LABEL", "Misura (SW 1K137)")

def enum_of(sig, val):
    e = dbc.get(sig)
    if not e or not e.get("values"): return None
    try: k = str(int(round(val)))
    except: return None
    return f"{k}={e['values'][k]}" if k in e["values"] else None

CATEG_OVERRIDE = {"IBATT1": "Scala/unità errata", "BAT_RINT_LOW_CUR_TRAITE": "Scala/unità errata",
                  "ATMO_PRESSURE_HIGH_ALTI": "CAN congelato (SW varia)", "PRESSION_ATMO": "CAN congelato (SW varia)"}
SENT = (254, 255, 65535, 2147483647, 214748364.7, 3276.7, 3276.8, 25.5, 25.4)
ACTIONABLE = {"SW non evolve (match nullo)", "Scala/unità errata", "Booleano/enum costante opposto"}

def categoria(r):
    s = r["status"]; sig = r["Signal Name"]; up = sig.upper()
    if s == "REMAP": return "Si insegue con rimappatura/sfasamento"
    if s == "CONFIG_EXCLUDED": return "Frame escluso da config veicolo"
    if s == "ISO_1K136": return "Comportamento ISO 1K136"
    if s == "SW_NOT_EVOLVING": return "SW non evolve (match nullo)"
    if s != "MISMATCH": return ""
    if sig in CATEG_OVERRIDE: return CATEG_OVERRIDE[sig]
    lo, hi = r.get("sig_min"), r.get("sig_max"); wlo, whi = r.get("sw_min"), r.get("sw_max")
    if "CHKSUM" in up or "CPT_PROCESS" in up: return "Checksum/contatore non allineato"
    if lo is not None and lo == hi:
        em = enum_of(sig, lo)
        if em and "nava" in em.lower(): return "CAN Unavailable"
    if wlo is not None and wlo == whi and any(abs(whi - x) < 1 for x in SENT): return "SW su valore invalido/init"
    if ("CPT" in up or "CNS_" in up) and lo == hi and wlo == whi: return "Contatore offset/sfasato"
    if lo == hi and wlo == whi and abs(hi) < 300 and abs(whi) < 300: return "Booleano/enum costante opposto"
    if r.get("kind") == "disc" and max(abs(lo), abs(hi), abs(wlo), abs(whi)) < 300: return "Booleano/enum costante opposto"
    return "Divergenza dinamica reale"

def verdict(r):
    s = r["status"]; sig = r["Signal Name"]
    if s == "OK":
        mf = r.get("match_frac"); return "OK - si inseguono", (f"match {mf*100:.0f}%" if mf is not None else "")
    if s == "REMAP": return "OK con rimappatura", "stesso andamento, codifica/fase diversa (es. enum invertita)"
    if s == "CONFIG_EXCLUDED":
        prev = {"NO_DATA_SIGNAL":"segnale non trasmesso","SW_NOT_EVOLVING":"SW non evolve","MISMATCH":"confronto non significativo","MISSING_SWVAR":"variabile SW assente"}.get(r.get("prev_status"),"segnale non acquisito")
        return "Giustificato (config veicolo)", f"Frame {r.get('frame_hex')} escluso da configurazione veicolo (Ctrl=FALSE): {prev} → giustificato"
    if s == "ISO_1K136":
        prev={"MISMATCH":"divergenza","SW_NOT_EVOLVING":"SW non evolve"}.get(r.get("prev_status_iso"),"deviazione")
        return "Giustificato (Comportamento ISO 1K136)", f"Deviazione ({prev}) presente in modo simile anche con la SW precedente 1K136 (check di trasparenza) -> non regressione del 1K137, giustificato"
    if s == "SW_NOT_EVOLVING":
        extra=" [deviazione presente SOLO in 1K137: in 1K136 il segnale insegue]" if r.get("only_1K137") else ""
        return "KO - SW non evolve", f"SW ferma a {r.get('sw_min')} mentre il CAN evolve [{r.get('sig_min')},{r.get('sig_max')}] → match nullo (regola 1)"+extra
    if s == "SUSPECT":
        mf = r.get("match_frac"); cyc = r.get("cycle_ms")
        if mf is not None:
            b = f"match {mf*100:.0f}% con tolleranza; contatore/checksum al limite soglia" if (r.get("kind")=="disc") else f"match {mf*100:.0f}%"
        else: b = "correlazione borderline"
        return "DA VERIFICARE", b
    if s == "MISMATCH":
        parts = []; lo, hi = r.get("sig_min"), r.get("sig_max"); wlo, whi = r.get("sw_min"), r.get("sw_max")
        if lo is not None and lo == hi:
            em = enum_of(sig, lo)
            if em: parts.append(f"CAN={em}")
        SM = {254:"254=NotAvailable",255:"255=NotAvailable",65535:"65535=NotAvailable",2147483647:"max int (init)",214748364.7:"max (init)",3276.7:"max (init)"}
        if wlo is not None and wlo == whi:
            for k,l in SM.items():
                if abs(whi-k)<1: parts.append(f"SW bloccata invalido/init ({l})"); break
        up = sig.upper()
        if ("CHKSUM" in up or "CPT_PROCESS" in up) and not parts:
            mf = r.get("match_frac"); parts.append(f"checksum/contatore non allineato (match {mf*100:.0f}%)" if mf is not None else "checksum/contatore non allineato")
        if not parts: parts.append("divergenza reale dei valori")
        return "KO - non si inseguono", "; ".join(parts)
    if s == "MISSING_SIGNAL": return "N/D - Signal assente", "canale CAN non presente nella misura"
    if s == "MISSING_SWVAR": return "N/D - Sw Variable assente", "variabile SW non presente nella misura"
    if s == "NO_DATA_SIGNAL": return "N/D - Signal senza campioni", "canale CAN presente ma mai trasmesso"
    if s == "NO_DATA_SWVAR": return "N/D - Sw senza campioni", "variabile SW presente ma senza campioni"
    if s in ("NO_OVERLAP","NO_VALID_DATA","MISSING_BOTH","NO_DATA_BOTH"): return "N/D", "dati non confrontabili"
    return s, ""

def fill_for(e):
    if e.startswith("Giustificato (Comportamento ISO"): return "D9EAD3"
    if e.startswith("Giustificato"): return "D6E4F0"
    if e.startswith("OK con"): return "B7E1CD"
    if e.startswith("OK"): return "C6EFCE"
    if e.startswith("DA VER"): return "FFEB9C"
    if e.startswith("KO"): return "FFC7CE"
    return "D9D9D9"

def fmt(v): return round(v,4) if isinstance(v,(int,float)) else ""

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Verifica coppie"
hdr = ["Signal Name (CAN)","Sw Variable","Esito","Categoria","Note","Bus","Messaggio","ECU tx","CMM dir","Cycle [ms]","Frame cfg","Unità","Match/Corr","Errore norm.","N camp.","sig min","sig max","sw min","sw max"]
ws.append(hdr)
for c in range(1,len(hdr)+1):
    cell=ws.cell(1,c); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="4472C4"); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
for r in results:
    es,note = verdict(r); sig=r["Signal Name"]; e=dbc.get(sig,{}); cat=categoria(r)
    if r.get("resolved_signal"): note=(note+"; " if note else "")+f"risolto su msg {r['resolved_signal'].split('.',1)[-1]}"
    mf=r.get("match_frac"); corr=r.get("corr")
    mc = (f"{r.get('remap_frac',0)*100:.0f}% remap" if r.get("status")=="REMAP" else f"{mf*100:.0f}% match") if mf is not None else (f"corr {corr:.2f}" if corr is not None else "")
    ws.append([sig,r["Sw Variable"],es,cat,note,e.get("bus",""),e.get("msg",""),e.get("tx",""),r.get("cmm_dir",""),r.get("cycle_ms",""),r.get("frame_cfg",""),e.get("unit",""),mc,fmt(r.get("norm_err")),r.get("n_samples",""),fmt(r.get("sig_min")),fmt(r.get("sig_max")),fmt(r.get("sw_min")),fmt(r.get("sw_max"))])
    ws.cell(ws.max_row,3).fill=PatternFill("solid",fgColor=fill_for(es))
    if cat in ACTIONABLE:
        ws.cell(ws.max_row,4).fill=PatternFill("solid",fgColor="ED7D31"); ws.cell(ws.max_row,4).font=Font(bold=True,color="FFFFFF")
for i,w in enumerate([30,26,24,28,44,6,18,12,8,9,9,8,13,11,7,9,9,9,9],1):
    ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(hdr))}{ws.max_row}"

from collections import Counter
cnt = Counter(verdict(r)[0] for r in results)
ws2 = wb.create_sheet("Riepilogo"); ws2.append(["Esito","N. coppie"])
for c in range(1,3): ws2.cell(1,c).font=Font(bold=True,color="FFFFFF"); ws2.cell(1,c).fill=PatternFill("solid",fgColor="4472C4")
order=["OK - si inseguono","OK con rimappatura","Giustificato (config veicolo)","Giustificato (Comportamento ISO 1K136)","DA VERIFICARE","KO - non si inseguono","KO - SW non evolve"]
rest=sorted(k for k in cnt if k not in order)
for k in order+rest:
    if k in cnt:
        ws2.append([k,cnt[k]]); ws2.cell(ws2.max_row,1).fill=PatternFill("solid",fgColor=fill_for(k))
ws2.append(["TOTALE",sum(cnt.values())]); ws2.cell(ws2.max_row,1).font=Font(bold=True); ws2.cell(ws2.max_row,2).font=Font(bold=True)
ws2.append([]); ws2.append(["Misura",MEAS]); ws2.append(["Coppie",len(results)])
ws2.column_dimensions["A"].width=34; ws2.column_dimensions["B"].width=30

# categoria breakdown
cc=Counter(categoria(r) for r in results if r["status"] in ("MISMATCH","SW_NOT_EVOLVING"))
ws2.append([]); ws2.append(["Categoria divergenza (KO)","N."])
hr=ws2.max_row
for c in range(1,3): ws2.cell(hr,c).font=Font(bold=True,color="FFFFFF"); ws2.cell(hr,c).fill=PatternFill("solid",fgColor="C0504D")
for k,v in cc.most_common(): ws2.append([k,v])

# Da correggere
ws3=wb.create_sheet("Da correggere (SW team)")
h3=["Priorità","Categoria","Signal Name","Sw Variable","Bus","Messaggio","CMM dir","Cycle [ms]","Unità","sig min","sig max","sw min","sw max","Note"]
ws3.append(h3)
for c in range(1,len(h3)+1): ws3.cell(1,c).font=Font(bold=True,color="FFFFFF"); ws3.cell(1,c).fill=PatternFill("solid",fgColor="ED7D31"); ws3.cell(1,c).alignment=Alignment(horizontal="center",wrap_text=True)
PRIO={"SW non evolve (match nullo)":1,"Scala/unità errata":1,"Booleano/enum costante opposto":2}
act=[r for r in results if categoria(r) in ACTIONABLE]; act.sort(key=lambda r:(PRIO.get(categoria(r),9),r["Signal Name"]))
for r in act:
    sig=r["Signal Name"]; e=dbc.get(sig,{}); cat=categoria(r); _,note=verdict(r)
    ws3.append([PRIO.get(cat,9),cat,sig,r["Sw Variable"],e.get("bus",""),e.get("msg",""),r.get("cmm_dir",""),r.get("cycle_ms",""),e.get("unit",""),fmt(r.get("sig_min")),fmt(r.get("sig_max")),fmt(r.get("sw_min")),fmt(r.get("sw_max")),note])
    col={"SW non evolve (match nullo)":"F4B6C2","Scala/unità errata":"FCE4D6","Booleano/enum costante opposto":"FFF2CC"}.get(cat,"FFF2CC")
    for c in range(1,len(h3)+1): ws3.cell(ws3.max_row,c).fill=PatternFill("solid",fgColor=col)
for i,w in enumerate([8,28,30,26,6,18,8,9,8,9,9,9,9,46],1): ws3.column_dimensions[get_column_letter(i)].width=w
ws3.freeze_panes="A2"

out=os.environ.get("OUT_XLSX", "Verifica_coppie.xlsx")
wb.save(out)
print("Saved:",out,"| azionabili:",len(act))
for k in order+rest:
    if k in cnt: print(f"  {k}: {cnt[k]}")
print("  TOT:",sum(cnt.values()))
